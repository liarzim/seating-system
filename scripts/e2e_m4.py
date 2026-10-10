"""Milestone 4 browser checks: org tree, data screen, Excel import, admin data tools.
Runs against http://localhost:5173 with the local test stack. Resets the local
admin password and reloads demo data first. Screenshots go to /tmp/claude-0/shots/."""
import os, re, subprocess, sys
from playwright.sync_api import sync_playwright, expect
import openpyxl

BASE = "http://localhost:5173"
SHOTS = "/tmp/claude-0/shots"
TMP = "/tmp/claude-0/m4"
os.makedirs(SHOTS, exist_ok=True); os.makedirs(TMP, exist_ok=True)
PW = "Admin-M4-2026x"

def sql(q):
    return subprocess.run(["psql", "-h", "/tmp", "-p", "55432", "-U", "postgres", "-d", "seating", "-tAc", q],
                          check=True, capture_output=True, text=True).stdout.strip()

sql("select public.admin_load_demo('דמה')")
sql(f"update auth.users set encrypted_password = crypt('{PW}', gen_salt('bf')) where email='admin@test.local'")
sql("update app_users set must_change_password=false, role='system_admin', active=true where email='admin@test.local'")
sql("delete from org_drafts")

def people_file(path, rows):
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append(["מספר עובד", "שם מלא", "אימייל", "יחידה", "תפקיד"])
    for r in rows: ws.append(r)
    wb.save(path)

unit = sql("select name from units where parent_id is not null order by name limit 1")
bad = f"{TMP}/people-bad.xlsx"
people_file(bad, [["90001", "בדיקה אחת", "", unit, "בודק"], ["90002", "בדיקה שתיים", "", "יחידה שלא קיימת", ""]])
good = f"{TMP}/people-good.xlsx"
people_file(good, [[str(90000 + i), f"עובד ייבוא {i}", f"imp{i}@example.com", unit, "בודק"] for i in range(1, 251)])

results = []
def step(name, fn):
    try:
        fn(); results.append(("PASS", name, ""))
    except Exception as e:  # noqa
        results.append(("FAIL", name, str(e).splitlines()[0][:220]))

with sync_playwright() as p:
    b = p.chromium.launch()
    ctx = b.new_context(viewport={"width": 1280, "height": 900}, locale="he-IL", accept_downloads=True)
    page = ctx.new_page()
    page.on("dialog", lambda d: d.accept())
    errors = []
    page.on("console", lambda m: errors.append(m.text) if m.type == "error" else None)

    def login():
        page.goto(BASE + "/login")
        page.get_by_label("אימייל").fill("admin@test.local")
        page.get_by_label("סיסמה", exact=True).fill(PW)
        page.get_by_role("button", name="כניסה", exact=True).click()
        expect(page).to_have_url(BASE + "/", timeout=10000)
        expect(page.get_by_role("link", name="מבנה ארגוני")).to_be_visible()
    step("admin logs in", login)

    def tree_loads():
        page.get_by_role("link", name="מבנה ארגוני").click()
        expect(page.locator(".tree > .t-unit")).to_have_count(7, timeout=10000)  # 6 divisions + "ללא שיוך"
        counts = page.locator(".tree > .t-unit > .t-row .t-count").all_inner_texts()
        assert sum(int(c) for c in counts) == 250, counts
        page.screenshot(path=f"{SHOTS}/m4-01-tree.png", full_page=True)
    step("tree shows 6 divisions + unassigned, 250 people in total", tree_loads)

    target = {}
    def search_opens():
        name = sql("select full_name from people where unit_id is not null order by employee_no limit 1")
        target["name"] = name
        target["from"] = sql(f"select u.name from people p join units u on u.id=p.unit_id where p.full_name='{name}' limit 1")
        page.get_by_label("חיפוש").fill(name)
        row = page.locator(".t-person.is-hit", has_text=name).first
        expect(row).to_be_visible(timeout=5000)
        page.screenshot(path=f"{SHOTS}/m4-02-search.png")
        page.get_by_label("חיפוש").fill("")
    step("search finds a person and opens the branch", search_opens)

    def move_now_and_undo():
        name = target["name"]
        page.get_by_label("חיפוש").fill(name)
        row = page.locator(".t-person", has_text=name).first
        row.hover(); row.get_by_role("button", name="העברה").click()
        page.locator("dialog[open]").get_by_label("יחידת יעד").select_option(label="חטיבת כספים")
        page.locator("dialog[open]").get_by_role("button", name="המשך").click()
        dlg = page.locator("dialog[open]")
        expect(dlg).to_contain_text("חטיבת כספים")
        page.screenshot(path=f"{SHOTS}/m4-03-move-dialog.png")
        dlg.get_by_role("button", name="העברה").click()
        expect(page.get_by_role("status")).to_contain_text("הועבר", timeout=5000)
        assert sql(f"select u.name from people p join units u on u.id=p.unit_id where p.full_name='{name}'") == "חטיבת כספים"
        page.get_by_role("button", name="ביטול ההעברה").click()
        expect(page.get_by_role("button", name="ביטול ההעברה")).to_have_count(0, timeout=5000)
        assert sql(f"select u.name from people p join units u on u.id=p.unit_id where p.full_name='{name}'") == target["from"]
        page.get_by_label("חיפוש").fill("")
    step("move now asks first, applies, and undo restores", move_now_and_undo)

    def drag_to_draft():
        name = target["name"]
        page.get_by_label("חיפוש").fill(name)
        src = page.locator(".t-person", has_text=name).first
        dst = page.locator(".tree > .t-unit > .t-row", has_text="חטיבת שירות").first
        src.drag_to(dst)
        dlg = page.locator("dialog[open]")
        expect(dlg).to_contain_text("חטיבת שירות", timeout=5000)
        dlg.get_by_label("להוסיף לטיוטה").check()
        dlg.get_by_label("שם הטיוטה").fill("בדיקת טיוטה")
        dlg.get_by_label("תאריך תוקף").fill("2027-01-03")
        dlg.get_by_role("button", name="הוספה לטיוטה").click()
        expect(page.locator(".draft-panel")).to_contain_text(name, timeout=5000)
        # nothing changed for real yet
        assert sql(f"select u.name from people p join units u on u.id=p.unit_id where p.full_name='{name}'") == target["from"]
        page.screenshot(path=f"{SHOTS}/m4-04-draft.png", full_page=True)
    step("drag and drop into a new dated draft (not applied yet)", drag_to_draft)

    def multi_select_into_draft():
        page.get_by_label("חיפוש").fill("")
        page.locator(".t-toggle[aria-label='פתיחת ללא שיוך']").click()
        boxes = page.locator(".tree > .t-unit:last-child .t-person input[type=checkbox]")
        n = min(3, boxes.count())
        assert n >= 2, f"only {boxes.count()} unassigned people"
        for i in range(n): boxes.nth(i).check()
        expect(page.locator(".selbar")).to_contain_text(str(n))
        page.get_by_role("button", name="העברת המסומנים").click()
        page.locator("dialog[open]").get_by_label("יחידת יעד").select_option(label="חטיבת דאטה")
        page.locator("dialog[open]").get_by_role("button", name="המשך").click()
        dlg = page.locator("dialog[open]")
        expect(dlg).to_contain_text(f"{n} אנשים")
        expect(dlg.get_by_label("להוסיף לטיוטה")).to_be_checked()  # viewing a draft -> defaults to it
        dlg.get_by_role("button", name="הוספה לטיוטה").click()
        expect(page.locator(".draft-list > li")).to_have_count(1 + n, timeout=5000)
        target["n"] = n
    step("multi-select several people into the same draft", multi_select_into_draft)

    def remove_and_apply_draft():
        page.locator(".draft-list > li").first.get_by_role("button", name="הסרה מהטיוטה").click()
        expect(page.locator(".draft-list > li")).to_have_count(target["n"], timeout=5000)
        page.locator(".draft-panel").get_by_role("button", name=re.compile("להחיל")).click()
        expect(page.locator(".draft-panel")).to_have_count(0, timeout=8000)
        cnt = sql("select count(*) from people p join units u on u.id=p.unit_id where u.name='חטיבת דאטה' and p.full_name not like 'x%'")
        assert sql("select status from org_drafts where name='בדיקת טיוטה'") == "applied"
        assert sql(f"select u.name from people p join units u on u.id=p.unit_id where p.full_name='{target['name']}'") == target["from"], "removed change was applied"
    step("remove one change, apply the draft now", remove_and_apply_draft)

    def unit_crud():
        page.get_by_role("button", name="יחידה חדשה").click()
        dlg = page.locator("dialog[open]")
        dlg.get_by_label("שם").fill("יחידת בדיקה")
        dlg.get_by_role("button", name="שמירה").click()
        row = page.locator(".t-row", has_text="יחידת בדיקה")
        expect(row).to_be_visible(timeout=5000)
        row.hover(); row.get_by_role("button", name="מחיקה").click()
        expect(page.locator(".t-row", has_text="יחידת בדיקה")).to_have_count(0, timeout=5000)
        # a unit with people cannot be deleted
        row = page.locator(".tree > .t-unit > .t-row", has_text="חטיבת כספים")
        row.hover(); row.get_by_role("button", name="מחיקה").click()
        expect(page.get_by_role("alert")).to_contain_text("אי אפשר למחוק", timeout=5000)
    step("add and delete a unit; non-empty unit is refused", unit_crud)

    def data_people():
        page.get_by_role("link", name="נתונים").click()
        expect(page.locator("table tbody tr").first).to_be_visible(timeout=10000)
        page.get_by_label("חיפוש").fill(target["name"])
        expect(page.locator("table tbody tr")).to_have_count(1)
        page.screenshot(path=f"{SHOTS}/m4-05-data-people.png", full_page=True)
    step("data screen lists and filters people", data_people)

    def secure_tab():
        page.get_by_role("tab", name="חדרים מאובטחים").click()
        expect(page.get_by_role("heading", name="אילו חדרים מאובטחים")).to_be_visible()
        page.screenshot(path=f"{SHOTS}/m4-06-secure.png", full_page=True)
    step("secure rooms tab renders", secure_tab)

    def import_bad_then_good():
        page.get_by_role("tab", name="ייבוא מאקסל").click()
        page.locator("input[type=file]").set_input_files(bad)
        expect(page.locator(".preview")).to_contain_text("שגיאות", timeout=10000)
        expect(page.locator(".preview table tbody tr")).to_have_count(1)
        assert sql("select count(*) from people where employee_no like '9000_'") == "0"
        page.screenshot(path=f"{SHOTS}/m4-07-import-errors.png", full_page=True)
        page.locator("input[type=file]").set_input_files(good)
        expect(page.locator(".preview")).to_contain_text("250 חדשים", timeout=15000)
        assert sql("select count(*) from people where employee_no::text ~ '^90[0-9]{3}$'") == "0", "dry run wrote rows"
        page.get_by_role("button", name="ייבוא 250 שורות").click()
        expect(page.get_by_role("status")).to_contain_text("250 נוספו", timeout=15000)
        assert sql("select count(*) from people") == "500"
        page.screenshot(path=f"{SHOTS}/m4-08-import-done.png", full_page=True)
    step("import: errors block everything, then 250 people preview and import", import_bad_then_good)

    def template_download():
        with page.expect_download() as d:
            page.get_by_role("button", name="הורדת הנתונים הקיימים").click()
        path = f"{TMP}/export.xlsx"; d.value.save_as(path)
        ws = openpyxl.load_workbook(path).worksheets[0]
        assert ws.max_row == 501, ws.max_row
        assert ws.sheet_view.rightToLeft
    step("export of existing people (500 rows, RTL)", template_download)

    def admin_reload_demo():
        page.get_by_role("tab", name="ניהול נתונים").click()
        btn = page.get_by_role("button", name="טעינת נתוני דמה")
        expect(btn).to_be_disabled()
        page.get_by_label(re.compile("כדי לאשר")).fill("דמה")
        btn.click()
        expect(page.get_by_role("status")).to_contain_text("נתוני הדמה נטענו", timeout=30000)
        assert sql("select count(*) from people") == "250"
    step("admin reloads demo data after typing the word", admin_reload_demo)

    def mobile():
        m = b.new_context(viewport={"width": 390, "height": 844}, locale="he-IL", color_scheme="dark", device_scale_factor=2)
        mp = m.new_page()
        mp.goto(BASE + "/login")
        mp.get_by_label("אימייל").fill("admin@test.local")
        mp.get_by_label("סיסמה", exact=True).fill(PW)
        mp.get_by_role("button", name="כניסה", exact=True).click()
        expect(mp).to_have_url(BASE + "/", timeout=10000)
        for path, shot in (("/org", "m4-09-org-mobile.png"), ("/data", "m4-10-data-mobile.png")):
            mp.goto(BASE + path)
            mp.wait_for_timeout(1500)
            w = mp.evaluate("document.documentElement.scrollWidth")
            mp.screenshot(path=f"{SHOTS}/{shot}", full_page=True)
            assert w <= 390, f"{path} horizontal scroll: {w}px"
        m.close()
    step("mobile + dark: org and data have no horizontal scroll", mobile)

    def visitor_blocked():
        v = b.new_context(locale="he-IL"); vp = v.new_page()
        vp.goto(BASE + "/data")
        expect(vp.locator(".notice")).to_contain_text("למנהלים", timeout=10000)
        assert vp.get_by_role("link", name="נתונים").count() == 0
        v.close()
    step("visitor cannot use the data screen", visitor_blocked)

    real = [e for e in errors if "Failed to load resource" not in e]
    step("no console errors", lambda: (_ for _ in ()).throw(AssertionError(real[:3])) if real else None)
    b.close()

for s, n, d in results:
    print(f"  {s}  {n}" + (f"  -> {d}" if d else ""))
failed = sum(1 for r in results if r[0] == "FAIL")
print(f"\n{len(results) - failed} passed, {failed} failed")
sys.exit(1 if failed else 0)
